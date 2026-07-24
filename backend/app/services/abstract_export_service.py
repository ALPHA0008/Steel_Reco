import io
import uuid

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from app.schemas.abstract import AbstractResponse
from app.services.abstract_service import AbstractService

_HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
_COMPUTED_FONT = Font(italic=True, color="1D4ED8")
_LABEL_FONT = Font(bold=True)
_TITLE_FONT = Font(bold=True, size=14)
_SUB_FONT = Font(size=10, italic=True, color="6B7280")


def _sum_by_dia(rows: list[dict], key: str) -> dict[str, float]:
    totals: dict[str, float] = {}
    for row in rows:
        dia = str(row["dia"])
        value = float(row.get(key) or 0)
        totals[dia] = totals.get(dia, 0.0) + value
    return totals


class AbstractExportService:
    """Renders the live Abstract into the same A-N x diameter grid the QS
    already knows from the legacy Excel -- something they can hand up the
    chain in the format leadership expects, generated from the ledger (never
    typed), same as the on-screen Abstract.
    """

    def __init__(self, session: AsyncSession) -> None:
        self._session = session

    async def export(self, project_id: uuid.UUID, project_name: str, year: int, month: int) -> bytes:
        abstract = await AbstractService(self._session).compute(project_id, year, month)
        return self._build_workbook(abstract, project_name)

    def _build_workbook(self, a: AbstractResponse, project_name: str) -> bytes:
        A = _sum_by_dia(a.section_a_received, "total_received_kg")
        B = _sum_by_dia(a.section_b_transferred, "total_transferred_kg")
        C = {k: float(v) for k, v in a.section_c_net_received.items()}
        D = _sum_by_dia(a.section_d_issued, "net_issued_kg")
        E = _sum_by_dia(a.section_e_consumption, "consumption_kg")
        F = _sum_by_dia(a.section_f_wip, "wip_kg")
        G = {k: float(v) for k, v in a.section_g_consumption_plus_wip.items()}
        H = {k: float(v) for k, v in a.section_h_theoretical_stock.items()}
        I_: dict[str, float] = {}
        J: dict[str, float] = {}
        for row in a.sections_ij_physical_stock:
            dia = str(row["dia"])
            I_[dia] = I_.get(dia, 0.0) + float(row.get("full_length_kg") or 0)
            J[dia] = J.get(dia, 0.0) + float(row.get("cut_piece_stock_kg") or 0)
        MyHome = {k: float(v) for k, v in a.section_myhome_stock.items()}
        K = {k: float(v) for k, v in a.section_k_total_physical.items()}
        L = {k: float(v) for k, v in a.section_l_wastage_qty.items()}
        m_pct = float(a.section_m_wastage_pct) if a.section_m_wastage_pct is not None else None
        n_kg = float(a.section_n_scrap_sold_kg)

        dias = sorted(
            {d for m in (A, B, C, D, E, F, G, H, I_, J, MyHome, K, L) for d in m},
            key=lambda x: float(x),
        )

        rows: list[tuple] = [
            ("A", "Received", A, False, None),
            ("B", "Transferred out", B, False, None),
            ("C", "Net Received", C, True, "= A - B"),
            ("C1", "Stock at My Home", MyHome, False, "steel at My Home's own yard"),
            ("D", "Issued to Contractor", D, False, "genuine sum, never = C"),
            ("E", "Consumption", E, False, None),
            ("F", "Work in Progress", F, False, None),
            ("G", "Consumption + WIP", G, True, "= E + F"),
            ("H", "Theoretical Stock", H, True, "= C - G"),
            ("I", "Physical — Full length", I_, False, None),
            ("J", "Physical — Cut pieces (stock)", J, False, None),
            ("K", "Total Physical", K, True, "= I + J + Stock at My Home"),
            ("L", "Wastage Qty", L, True, "= H - K"),
        ]

        wb = Workbook()
        ws = wb.active
        ws.title = "Abstract"

        ws.merge_cells("A1:C1")
        ws["A1"] = f"{project_name} — Monthly Steel Abstract"
        ws["A1"].font = _TITLE_FONT
        ws.merge_cells("A2:C2")
        ws["A2"] = f"Period: {a.period_label} (cumulative since project start) — {a.pipeline_version}"
        ws["A2"].font = _SUB_FONT

        header_row = 4
        ws.cell(header_row, 1, "Section").font = _HEADER_FONT
        ws.cell(header_row, 1).fill = _HEADER_FILL
        for i, dia in enumerate(dias):
            c = ws.cell(header_row, 2 + i, f"{float(dia):g} mm")
            c.font = _HEADER_FONT
            c.fill = _HEADER_FILL
            c.alignment = Alignment(horizontal="right")
        total_col = 2 + len(dias)
        tc = ws.cell(header_row, total_col, "Total (MT)")
        tc.font = _HEADER_FONT
        tc.fill = _HEADER_FILL
        tc.alignment = Alignment(horizontal="right")

        r = header_row + 1
        for code, label, values, computed, formula in rows:
            label_cell = ws.cell(r, 1, f"{code} · {label}" + (f"  [{formula}]" if formula else ""))
            label_cell.font = _COMPUTED_FONT if computed else _LABEL_FONT
            total_mt = 0.0
            for i, dia in enumerate(dias):
                kg = values.get(dia, 0.0)
                mt = kg / 1000
                total_mt += mt
                cell = ws.cell(r, 2 + i, round(mt, 2) if kg else None)
                cell.alignment = Alignment(horizontal="right")
            tcell = ws.cell(r, total_col, round(total_mt, 2))
            tcell.font = Font(bold=True)
            tcell.alignment = Alignment(horizontal="right")
            r += 1

        m_cell_label = ws.cell(r, 1, "M · Wastage %  [= L / G]")
        m_cell_label.font = _COMPUTED_FONT
        ws.cell(r, total_col, round(m_pct, 2) if m_pct is not None else None).alignment = Alignment(horizontal="right")
        r += 1

        ws.cell(r, 1, "N · Scrap Sold").font = _LABEL_FONT
        ws.cell(r, total_col, round(n_kg / 1000, 2)).alignment = Alignment(horizontal="right")
        r += 2

        if a.findings:
            ws.cell(r, 1, "Findings").font = _LABEL_FONT
            r += 1
            for f in a.findings:
                ws.cell(r, 1, f"[{f['severity'].upper()}] {f['message']}").font = Font(size=9, color="B45309")
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=total_col)
                r += 1

        ws.column_dimensions["A"].width = 34
        for i in range(len(dias) + 1):
            ws.column_dimensions[get_column_letter(2 + i)].width = 13
        ws.freeze_panes = ws.cell(header_row + 1, 2).coordinate

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
