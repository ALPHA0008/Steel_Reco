"""Parser for the company-wide "Copy of Recon Steel 28-02-2026.xlsx" workbook.

This is the SINGLE source of truth for the four real sites we load into the
platform for the admin multi-site dashboard: Grava, Sayuk, Nishada, 99.

What the workbook actually is
-----------------------------
A multi-site monthly reconciliation book. Each monthly sheet (there are ~51,
from 'December 21' through 'February - 26') lays the sites out ACROSS columns:

  * Row 3  = the site-name header row. The column a given site occupies MOVES
             from sheet to sheet (sites were added over the years), so we read
             this header dynamically on every sheet -- never hardcode a column.
  * Rows 6,7,8,9,11,13,15,17,19,21,23,25,27,29 are the reconciliation lines,
    and this row->metric layout is STABLE across every era of the book:

        6  Total Received as per Stores Record
        7  Received from Other Sites
        8  Transfer to Other Sites
        9  A  Net Received Qty
        11 B  Total reinforcement steel issued to Contractors
        13 C  Consumption Details
        15 D  Work In Progress quantity
        17 E  TOTAL (Consumption + WIP  = C + D)
        19 F  Balance as per statement (A - E)
        21 G  Physical stock (Full Lengths)
        23 H  Cut Pieces Available
        25    Steel Scrap Sold in MT
        27 I  Wastage = Theoretical stock - Physical stock
        29 J  % Wastage  (= I / E)

All values are in METRIC TONNES (MT) in the sheet.

IMPORTANT: this book carries SUMMARY numbers only -- there are no underlying
GRN / issue / JMR transaction rows here (unlike the APAS SAP dump). So the
loader that consumes this parser loads month-end SNAPSHOT numbers as-is and
fills only the missing transaction-level *detail* with representative mock
rows -- see load_recon_sites.py.

This module has no DB or app dependency -- pure openpyxl -> plain dicts, so it
can be unit-run standalone (`python scripts/recon_sites_parser.py`).
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl

# The workbook lives at the repo root (one level above backend/).
DEFAULT_WORKBOOK = Path(__file__).resolve().parents[2] / "Copy of Recon Steel 28-02-2026.xlsx"

# The four sites we ingest. Matched EXACTLY against the row-3 header text
# (stripped). "Grava" is the parent site; "Grava Residences PC/CIS" are
# separate columns we deliberately do NOT fold in (they are their own sites in
# the book). "99" is a real site name in the book.
TARGET_SITES = ["Grava", "Sayuk", "Nishada", "99"]

# Sheets that are NOT month sheets (roll-ups / duplicates) -- skipped.
SKIP_SHEETS = {"Abstract", "February - 26 (2)"}

# row (1-indexed) -> metric key. Stable across all eras of the workbook.
_ROW_METRIC = {
    6: "total_received",
    7: "received_other_sites",
    8: "transfer_other_sites",
    9: "A_net_received",
    11: "B_issued",
    13: "C_consumption",
    15: "D_wip",
    17: "E_total",
    19: "F_balance",
    21: "G_physical_full",
    23: "H_cut_pieces",
    25: "scrap_sold",
    27: "I_wastage_qty",
    29: "J_wastage_pct",
}

_MONTHS_FULL = {
    "january": 1, "february": 2, "march": 3, "april": 4, "may": 5, "june": 6,
    "july": 7, "august": 8, "september": 9, "october": 10, "november": 11, "december": 12,
}
_MONTHS_SHORT = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}


def parse_sheet_month(sheet_name: str) -> tuple[int, int] | None:
    """'February - 26' / 'May 24' / 'Mar 21' / 'December 23' -> (year, month).

    Returns None for anything that is not a month sheet.
    """
    s = sheet_name.strip().lower()
    m = re.match(r"([a-z]+)\s*-?\s*'?(\d{2})\b", s)
    if not m:
        return None
    word, yy = m.group(1), int(m.group(2))
    month = _MONTHS_FULL.get(word) or _MONTHS_SHORT.get(word[:3])
    if not month:
        return None
    return (2000 + yy, month)


def _num(value) -> float | None:
    """Return a float only for genuine numbers; None for text like
    'Not Received' or blank cells."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return None


def parse_recon_sites(workbook_path: Path | str = DEFAULT_WORKBOOK) -> dict:
    """Parse every month sheet and return the full per-site history.

    Structure returned::

        {
          "Grava": {
             (2026, 2): {"A_net_received": 21651.07, "C_consumption": ...,
                         "D_wip": ..., "G_physical_full": ..., "H_cut_pieces": ...,
                         "scrap_sold": ..., "J_wastage_pct": 0.0389, ...},
             (2026, 1): {...},
             ...
          },
          "Sayuk": {...}, "Nishada": {...}, "99": {...},
        }

    Values are in MT (as the sheet stores them). Only genuinely-numeric cells
    are captured; text sentinels ('Not Received') become absent keys.
    """
    path = Path(workbook_path)
    if not path.exists():
        raise FileNotFoundError(f"Recon Steel workbook not found at {path}")

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out: dict[str, dict[tuple[int, int], dict[str, float]]] = {s: {} for s in TARGET_SITES}
    sheets_parsed = 0

    for name in wb.sheetnames:
        if name in SKIP_SHEETS:
            continue
        ym = parse_sheet_month(name)
        if ym is None:
            continue

        ws = wb[name]
        grid = [list(r) for r in ws.iter_rows(values_only=True, max_row=30, max_col=32)]
        if len(grid) < 9:
            continue

        header = grid[2]  # row 3
        col_of_site: dict[str, int] = {}
        for c, v in enumerate(header):
            if v is None:
                continue
            label = str(v).strip()
            if label in TARGET_SITES:
                col_of_site[label] = c
        if not col_of_site:
            continue

        sheets_parsed += 1
        for site, col in col_of_site.items():
            record: dict[str, float] = {}
            for row_idx, metric in _ROW_METRIC.items():
                row = grid[row_idx - 1] if row_idx - 1 < len(grid) else None
                cell = row[col] if row and col < len(row) else None
                val = _num(cell)
                if val is not None:
                    record[metric] = val
            if record:
                out[site][ym] = record

    wb.close()
    out["_meta"] = {"sheets_parsed": sheets_parsed, "workbook": str(path)}
    return out


if __name__ == "__main__":
    data = parse_recon_sites()
    meta = data.pop("_meta")
    print(f"Parsed {meta['sheets_parsed']} month sheets from {meta['workbook']}\n")
    for site in TARGET_SITES:
        months = sorted(data[site])
        if not months:
            print(f"{site}: NO DATA")
            continue
        first, last = months[0], months[-1]
        rec = data[site][last]
        print(f"{site:9s}  {len(months):2d} months  {first}..{last}")
        print("   latest:", ", ".join(f"{k}={v:.2f}" for k, v in rec.items()))
