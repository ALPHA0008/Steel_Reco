"""Sections I (Physical full-length) and J (Physical cut pieces) — parsed
independently from Annexure-1 / Annexure-2, which repeat one block per
contractor (KLC, then GLC) in the same sheet.
"""

from pathlib import Path

import openpyxl

from .dia import DIAS, empty_dia_totals, parse_dia


def _parse_annexure_blocks(rows: list[tuple], dia_col: int, weight_col: int) -> dict[int, float]:
    """Walk every row; whenever the Dia cell parses to a known diameter, add
    its weight cell. Works across an arbitrary number of per-contractor
    blocks since we don't rely on block boundaries, only on row shape.
    """
    totals = empty_dia_totals()
    for row in rows:
        if dia_col >= len(row):
            continue
        dia = parse_dia(row[dia_col] if isinstance(row[dia_col], str) else None)
        if dia is None:
            continue
        weight = row[weight_col] if weight_col < len(row) else None
        if weight is None:
            continue
        totals[dia] += float(weight)
    return totals


def parse_section_i(workbook_path: Path) -> dict[int, float]:
    """Physical stock, full lengths (bundles + loose rods) — Annexure-1, col index 1=Dia, 8=Grand Total Weight."""
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    ws = wb["Annexure-1"]
    rows = list(ws.iter_rows(values_only=True))
    return _parse_annexure_blocks(rows, dia_col=1, weight_col=8)


def parse_section_j(workbook_path: Path) -> dict[int, float]:
    """Physical stock, cut pieces — Annexure-2, col index 1=Dia, 4=Grand Total Weight."""
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    ws = wb["Annexure-2"]
    rows = list(ws.iter_rows(values_only=True))
    return _parse_annexure_blocks(rows, dia_col=1, weight_col=4)


if __name__ == "__main__":
    import sys

    wb_path = Path(sys.argv[1])
    i_totals = parse_section_i(wb_path)
    j_totals = parse_section_j(wb_path)
    print(f"Section I (Physical full): {i_totals}  TOTAL={sum(i_totals.values()):.3f}")
    print(f"Section J (Physical cut):  {j_totals}  TOTAL={sum(j_totals.values()):.3f}")
